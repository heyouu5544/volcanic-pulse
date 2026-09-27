# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///

from __future__ import annotations

import json
from collections import defaultdict
from html import escape
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PROJECT_ROOT = HERE.parent
DATA_FILE = PROJECT_ROOT / "data" / "GVP_Eruption_List_Holocene_20260424.xlsx"
OUTPUT_FILE = HERE / "index.html"

YEAR_MIN = 2000
YEAR_MAX = 2025
MISSING_VEI = {"", "NULL", "null", "None"}
FILTER_ORDER = ["low", "medium", "high", "unknown"]
FILTER_LABELS = {
    "low": "Low / 0–1",
    "medium": "Medium / 2–3",
    "high": "High / 4+",
    "unknown": "Unknown",
}
FILTER_COLORS = {
    "low": "#F5C76E",
    "medium": "#FF8C42",
    "high": "#FF593B",
    "unknown": "#7E8AA3",
}


def parse_vei(value):
    if value is None or str(value).strip() in MISSING_VEI:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def vei_group(value):
    if value is None:
        return "unknown"
    if value <= 1:
        return "low"
    if value <= 3:
        return "medium"
    return "high"


def start_date_string(year, month, day):
    pieces = []
    if month is not None:
        try:
            month_num = int(month)
        except (TypeError, ValueError):
            month_num = None
        if month_num is not None:
            pieces.append(f"{month_num:02d}")
    if day is not None:
        try:
            day_num = int(day)
        except (TypeError, ValueError):
            day_num = None
        if day_num is not None:
            pieces.append(f"{day_num:02d}")
    if not pieces:
        return str(year)
    if len(pieces) == 1:
        return f"{year}-{pieces[0]}"
    return f"{year}-{pieces[0]}-{pieces[1]}"


def load_eruption_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Eruption List"]
    header_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    header_map = {name: idx for idx, name in enumerate(header_row)}

    required = [
        "Volcano Name",
        "Start Year",
        "Start Month",
        "Start Day",
        "VEI",
    ]
    missing = [name for name in required if name not in header_map]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    rows = []
    year_counts = defaultdict(int)
    parse_errors = 0
    for row in sheet.iter_rows(min_row=3, values_only=True):
        try:
            year_value = row[header_map["Start Year"]]
            if year_value is None:
                continue
            year = int(year_value)
            if not (YEAR_MIN <= year <= YEAR_MAX):
                continue

            volcano = row[header_map["Volcano Name"]]
            raw_vei = row[header_map["VEI"]]
            vei = parse_vei(raw_vei)
            year_counts[year] += 1
            eruption_number = row[header_map["Eruption Number"]] if "Eruption Number" in header_map else None
            item = {
                "id": f"{year}-{year_counts[year]}-{volcano or 'unknown'}",
                "volcano": str(volcano or "Unknown volcano"),
                "year": year,
                "month": row[header_map["Start Month"]],
                "day": row[header_map["Start Day"]],
                "vei": vei,
                "vei_label": "Unknown" if vei is None else f"{vei:g}",
                "group": vei_group(vei),
                "order": year_counts[year],
                "date": start_date_string(year, row[header_map["Start Month"]], row[header_map["Start Day"]]),
                "eruption_number": None if eruption_number in (None, "") else str(eruption_number),
            }
            rows.append(item)
        except Exception:
            parse_errors += 1

    workbook.close()
    return rows, parse_errors


def build_svg_background():
    constellations = []
    for i in range(90):
        x = 80 + ((i * 97) % 1040)
        y = 80 + ((i * 131) % 520)
        size = 1.2 + (i % 3) * 0.7
        alpha = 0.25 + (i % 6) * 0.1
        constellations.append(f'<circle cx="{x}" cy="{y}" r="{size}" fill="rgba(255,255,255,{alpha})" />')

    rings = []
    for radius in (170, 230, 300, 370):
        rings.append(f'<circle cx="610" cy="420" r="{radius}" fill="none" stroke="rgba(56,214,230,0.22)" stroke-width="1" />')
    for angle in range(0, 360, 24):
        radians = angle * 3.141592653589793 / 180
        x1 = 610 + 120 * __import__("math").cos(radians)
        y1 = 420 + 120 * __import__("math").sin(radians)
        x2 = 610 + 430 * __import__("math").cos(radians)
        y2 = 420 + 430 * __import__("math").sin(radians)
        rings.append(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="rgba(56,214,230,0.16)" stroke-width="1" />')
    return "\n".join(constellations + rings)


def build_page(rows, parse_errors):
    unknown_count = sum(1 for item in rows if item["vei"] is None)
    total_count = len(rows)
    y_counts = defaultdict(int)
    for item in rows:
        y_counts[item["year"]] += 1
    max_year_count = max(y_counts.values()) if y_counts else 0

    # JS must run without bundling; it reads the embedded data below.
    dataset_json = json.dumps(rows, ensure_ascii=False)
    page = f'''<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>Volcanic Pulse — Interactive</title>
  <style>
    :root {{
      --bg: #090B10;
      --bg-2: #111722;
      --panel: rgba(15, 18, 25, 0.78);
      --panel-strong: rgba(18, 22, 30, 0.92);
      --gold: #F5C76E;
      --gold-soft: rgba(245, 199, 110, 0.18);
      --amber: #FF8C42;
      --lava: #FF593B;
      --cyan: #38D6E6;
      --cyan-soft: rgba(56, 214, 230, 0.18);
      --violet: #C58CF7;
      --steel: #8F9FB6;
      --text: #E8EEF4;
      --muted: #96A6BC;
      --line: rgba(255,255,255,0.13);
      --shadow: rgba(10, 12, 20, 0.8);
    }}

    * {{ box-sizing: border-box; }}

    html, body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: "Segoe UI", Arial, sans-serif;
      min-height: 100%;
    }}

    body {{
      background:
        radial-gradient(circle at 50% 15%, rgba(245, 199, 110, 0.12), transparent 30%),
        radial-gradient(circle at 80% 18%, rgba(56, 214, 230, 0.1), transparent 28%),
        var(--bg);
      overflow-x: hidden;
    }}

    .shell {{
      width: min(1500px, 96vw);
      margin: 24px auto 48px;
      padding: 18px 18px 10px;
      border: 1px solid var(--line);
      background: rgba(11, 14, 20, 0.82);
      box-shadow: 0 0 0 1px rgba(255,255,255,0.02), 0 30px 70px rgba(0,0,0,0.45);
      position: relative;
    }}

    .header-bar {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 14px;
      margin-bottom: 18px;
      padding: 6px 2px 12px;
      border-bottom: 1px solid var(--line);
    }}

    .brand {{
      white-space: nowrap;
      letter-spacing: 0.34em;
      font-weight: 800;
      font-size: 11px;
      color: var(--muted);
      text-transform: uppercase;
    }}

    .status {{
      font-size: 11px;
      color: var(--muted);
      letter-spacing: 0.12em;
      text-transform: uppercase;
      border: 1px solid var(--line);
      background: rgba(255,255,255,0.02);
      padding: 8px 12px;
    }}

    .title-row {{
      display: grid;
      grid-template-columns: minmax(0, 1.2fr) minmax(280px, 430px);
      gap: 22px;
      margin-bottom: 16px;
      align-items: end;
    }}

    .hero {{
      padding: 18px 0 0 10px;
      position: relative;
    }}

    .eyebrow {{
      display: inline-flex;
      align-items: center;
      gap: 10px;
      margin-bottom: 14px;
      color: var(--muted);
      font-size: 9px;
      letter-spacing: 0.26em;
      text-transform: uppercase;
    }}

    .eyebrow::before {{
      content: "";
      display: inline-block;
      width: 34px;
      height: 1px;
      background: linear-gradient(90deg, var(--gold), transparent);
    }}

    .title-stack {{
      display: flex;
      flex-direction: column;
      gap: 12px;
    }}

    .year-badge {{
      display: inline-flex;
      align-items: baseline;
      gap: 10px;
      align-self: flex-start;
      margin-top: 2px;
      padding: 8px 12px;
      border: 1px solid rgba(56,214,230,0.32);
      background: linear-gradient(180deg, rgba(56,214,230,0.08), rgba(17,26,32,0.12));
      box-shadow: inset 0 0 18px rgba(56,214,230,0.07), 0 0 18px rgba(56,214,230,0.08);
      backdrop-filter: blur(2px);
    }}

    .year-badge span {{
      color: var(--muted);
      font-size: 9px;
      letter-spacing: 0.2em;
      text-transform: uppercase;
    }}

    .year-badge strong {{
      color: var(--cyan);
      font-size: 0.92rem;
      letter-spacing: 0.12em;
      font-weight: 800;
    }}

    h1 {{
      margin: 0;
      display: flex;
      flex-direction: column;
      gap: 3px;
      font-size: clamp(2.7rem, 6vw, 7.2rem);
      line-height: 0.82;
      letter-spacing: 0.08em;
      text-transform: uppercase;
      position: relative;
      display: inline-block;
    }}

    h1 .word-volcanic {{
      display: block;
      font-weight: 300;
      letter-spacing: 0.18em;
      color: var(--text);
      text-shadow: 0 0 20px rgba(255,255,255,0.08);
    }}

    h1 .word-pulse {{
      display: block;
      font-weight: 900;
      letter-spacing: 0.08em;
      color: var(--gold);
      text-shadow: 0 0 28px rgba(245, 199, 110, 0.28);
    }}

    h1::after {{
      content: "";
      position: absolute;
      left: 0;
      right: 6%;
      bottom: -12px;
      height: 1px;
      background: linear-gradient(90deg, var(--gold), rgba(245,199,110,0.18), transparent);
    }}

    .subtitle {{
      margin-top: 18px;
      color: var(--muted);
      font-size: 11px;
      letter-spacing: 0.2em;
      text-transform: uppercase;
      max-width: 690px;
      line-height: 1.7;
    }}

    .panel {{
      background: linear-gradient(180deg, rgba(17,20,29,0.9), rgba(12,16,22,0.78));
      border: 1px solid var(--line);
      box-shadow: inset 0 0 0 1px rgba(255,255,255,0.02), 0 0 0 1px rgba(56,214,230,0.04);
      padding: 16px 16px 14px;
      position: relative;
      overflow: hidden;
    }}

    .panel::before {{
      content: "";
      position: absolute;
      inset: 0;
      background: linear-gradient(120deg, transparent 0%, rgba(56,214,230,0.08), transparent 40%);
      pointer-events: none;
    }}

    .metrics {{
      display: grid;
      grid-template-columns: repeat(2, minmax(120px, 1fr));
      gap: 12px;
      margin-top: 4px;
    }}

    .metric {{
      border: 1px solid rgba(255,255,255,0.08);
      background: linear-gradient(180deg, rgba(255,255,255,0.03), rgba(10,13,20,0.2));
      padding: 12px 10px;
      min-height: 90px;
      box-shadow: inset 0 0 0 1px rgba(255,255,255,0.015);
    }}

    .metric-label {{
      display: block;
      color: var(--muted);
      letter-spacing: 0.18em;
      text-transform: uppercase;
      font-size: 9px;
      margin-bottom: 10px;
    }}

    .metric-value {{
      font-size: clamp(1.5rem, 2.7vw, 2.4rem);
      font-weight: 800;
      letter-spacing: 0.02em;
      color: var(--text);
      text-shadow: 0 0 14px rgba(255,255,255,0.08);
    }}

    .metric-value.gold {{ color: var(--gold); }}
    .metric-value.cyan {{ color: var(--cyan); }}
    .metric-value.orange {{ color: var(--amber); }}
    .metric-value.red {{ color: var(--lava); }}

    .chart-wrap {{
      position: relative;
      border: 1px solid var(--line);
      background: linear-gradient(180deg, rgba(10,12,18,0.82), rgba(9,11,16,0.96));
      min-height: 0;
      overflow: hidden;
      box-shadow: inset 0 0 25px rgba(56,214,230,0.06), 0 0 22px rgba(56,214,230,0.03);
    }}

    .chart-wrap::before {{
      content: "";
      position: absolute;
      inset: 12px 12px auto 12px;
      height: 1px;
      background: linear-gradient(90deg, rgba(245,199,110,0.8), rgba(56,214,230,0.5), transparent 90%);
      pointer-events: none;
    }}

    .chart-wrap::after {{
      content: "";
      position: absolute;
      inset: 0;
      background: radial-gradient(circle at center, transparent 52%, rgba(56,214,230,0.04), transparent 74%);
      pointer-events: none;
    }}

    .hud {{
      position: absolute;
      inset: 0;
      pointer-events: none;
    }}

    .hud .corner {{
      position: absolute;
      width: 32px;
      height: 32px;
      border-color: rgba(56,214,230,0.7);
      border-style: solid;
      pointer-events: none;
    }}

    .hud .corner.tl {{ top: 14px; left: 14px; border-width: 1px 0 0 1px; }}
    .hud .corner.tr {{ top: 14px; right: 14px; border-width: 1px 1px 0 0; }}
    .hud .corner.bl {{ bottom: 14px; left: 14px; border-width: 0 0 1px 1px; }}
    .hud .corner.br {{ bottom: 14px; right: 14px; border-width: 0 1px 1px 0; }}

    svg {{
      display: block;
      width: 100%;
      height: auto;
      aspect-ratio: 1200 / 980;
      background: transparent;
    }}

    .top-year-control {{
      position: relative;
      z-index: 12;
      margin: 0;
      padding: 12px 14px 10px;
      background: linear-gradient(180deg, rgba(17,20,28,0.82), rgba(13,16,22,0.76));
      border-bottom: 1px solid var(--line);
      box-shadow: inset 0 0 0 1px rgba(255,255,255,0.02);
    }}

    .top-year-control .slider-wrap {{
      margin: 0;
    }}

    .controls {{
      display: grid;
      grid-template-columns: minmax(0, 1.1fr) minmax(300px, 0.9fr);
      gap: 18px;
      margin-top: 18px;
    }}

    .control-panel {{
      padding: 14px 16px 12px;
      background: rgba(12,15,20,0.88);
      border: 1px solid var(--line);
      box-shadow: inset 0 1px 0 rgba(255,255,255,0.03);
    }}

    .control-head {{
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 10px;
      margin-bottom: 12px;
      color: var(--muted);
      font-size: 11px;
      letter-spacing: 0.18em;
      text-transform: uppercase;
    }}

    .play-row {{
      display: flex;
      align-items: center;
      flex-wrap: wrap;
      gap: 10px;
      margin-bottom: 12px;
    }}

    button, select, input[type="range"] {{
      font: inherit;
    }}

    .btn {{
      border: 1px solid rgba(255,255,255,0.08);
      background: rgba(255,255,255,0.02);
      color: var(--text);
      letter-spacing: 0.1em;
      text-transform: uppercase;
      font-size: 10px;
      padding: 10px 14px;
      cursor: pointer;
      transition: transform 0.2s ease, border-color 0.2s ease, box-shadow 0.2s ease;
      box-shadow: inset 0 0 0 1px rgba(255,255,255,0.02);
    }}

    .btn:hover {{
      transform: translateY(-1px);
      border-color: rgba(56,214,230,0.5);
      box-shadow: 0 0 0 1px rgba(56,214,230,0.12);
    }}

    .btn.primary {{
      background: linear-gradient(180deg, rgba(245,199,110,0.18), rgba(245,199,110,0.05));
      border-color: rgba(245,199,110,0.6);
      color: var(--gold);
      box-shadow: 0 0 18px rgba(245,199,110,0.08);
    }}

    .speed-group {{
      display: inline-flex;
      gap: 6px;
      background: rgba(255,255,255,0.03);
      border: 1px solid var(--line);
      padding: 5px;
    }}

    .speed-group button {{
      background: transparent;
      border: 1px solid transparent;
      color: var(--muted);
      padding: 8px 10px;
      min-width: 58px;
      font-size: 10px;
      letter-spacing: 0.12em;
      text-transform: uppercase;
      cursor: pointer;
    }}

    .speed-group button.active {{
      border-color: rgba(56,214,230,0.55);
      background: rgba(56,214,230,0.08);
      color: var(--text);
    }}

    .slider-wrap {{
      display: grid;
      grid-template-columns: 58px minmax(0, 1fr) 62px;
      gap: 12px;
      align-items: center;
      margin-bottom: 14px;
    }}

    .slider-label {{
      color: var(--muted);
      font-size: 10px;
      letter-spacing: 0.14em;
      text-transform: uppercase;
    }}

    input[type="range"] {{
      width: 100%;
      accent-color: var(--gold);
    }}

    .year-num {{
      text-align: right;
      color: var(--gold);
      font-weight: 700;
      font-size: 1.1rem;
    }}

    .mode-row {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      margin-top: 8px;
    }}

    .toggle {{
      position: relative;
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 10px;
      cursor: pointer;
      border: 1px solid var(--line);
      background: rgba(255,255,255,0.02);
      color: var(--muted);
      font-size: 10px;
      letter-spacing: 0.12em;
      text-transform: uppercase;
    }}

    .toggle input {{
      accent-color: var(--gold);
      width: 14px;
      height: 14px;
      margin: 0;
    }}

    .filters {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      margin-top: 8px;
    }}

    .filter-chip {{
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 10px;
      border: 1px solid var(--line);
      background: rgba(255,255,255,0.02);
      font-size: 10px;
      letter-spacing: 0.12em;
      text-transform: uppercase;
      color: var(--muted);
    }}

    .filter-chip .dot {{
      width: 10px;
      height: 10px;
      border-radius: 50%;
      display: inline-block;
      border: 1px solid rgba(255,255,255,0.5);
      box-shadow: 0 0 12px rgba(255,255,255,0.15);
    }}

    .filter-chip input {{
      accent-color: var(--gold);
      width: 14px;
      height: 14px;
      margin: 0;
    }}

    .side-panel {{
      padding: 14px 16px 10px;
      background: rgba(12,15,20,0.88);
      border: 1px solid var(--line);
      box-shadow: inset 0 1px 0 rgba(255,255,255,0.03);
    }}

    .detail-header {{
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 10px;
      padding-bottom: 8px;
      border-bottom: 1px solid var(--line);
      margin-bottom: 10px;
      color: var(--muted);
      font-size: 10px;
      letter-spacing: 0.18em;
      text-transform: uppercase;
    }}

    .detail-body {{
      min-height: 146px;
    }}

    .detail-name {{
      margin: 6px 0 8px;
      font-weight: 800;
      font-size: clamp(1.2rem, 2vw, 2rem);
      line-height: 1.1;
      color: var(--text);
    }}

    .detail-list {{
      list-style: none;
      padding: 0;
      margin: 0;
      display: grid;
      gap: 8px;
      color: var(--muted);
      font-size: 12px;
      letter-spacing: 0.06em;
    }}

    .detail-list strong {{
      color: var(--text);
      font-weight: 700;
    }}

    .detail-chip {{
      display: inline-flex;
      align-items: center;
      gap: 8px;
      margin-top: 12px;
      padding: 8px 10px;
      border: 1px solid rgba(255,255,255,0.12);
      font-size: 10px;
      letter-spacing: 0.12em;
      text-transform: uppercase;
      color: var(--muted);
      background: rgba(255,255,255,0.02);
    }}

    .detail-chip .dot {{
      width: 10px;
      height: 10px;
      border-radius: 50%;
      display: inline-block;
      background: var(--gold);
      box-shadow: 0 0 12px rgba(245,199,110,0.6);
    }}

    .tooltip {{
      position: absolute;
      pointer-events: none;
      background: rgba(7, 10, 18, 0.96);
      color: var(--text);
      border: 1px solid rgba(245, 199, 110, 0.8);
      box-shadow: 0 0 0 1px rgba(56,214,230,0.25), 0 18px 38px rgba(0,0,0,0.42);
      padding: 10px 12px 8px;
      font-size: 12px;
      line-height: 1.35;
      min-width: 170px;
      transform: translate(14px, -12px);
      opacity: 0;
      transition: opacity 0.18s ease;
      z-index: 30;
    }}

    .tooltip::before {{
      content: "";
      position: absolute;
      left: -12px;
      top: 14px;
      width: 12px;
      height: 1px;
      background: linear-gradient(90deg, var(--cyan), rgba(56,214,230,0.15));
      box-shadow: 0 0 10px rgba(56,214,230,0.35);
    }}

    .tooltip.visible {{ opacity: 1; }}

    .tooltip-head {{
      color: var(--gold);
      margin-bottom: 8px;
      font-size: 13px;
      font-weight: 800;
      letter-spacing: 0.08em;
      text-transform: uppercase;
    }}

    .tooltip-row {{
      display: grid;
      grid-template-columns: 52px 1fr;
      gap: 8px;
      font-size: 9px;
      letter-spacing: 0.14em;
      text-transform: uppercase;
      color: var(--muted);
      padding-top: 3px;
    }}

    .tooltip-row strong {{
      color: var(--text);
      font-weight: 700;
      letter-spacing: 0.04em;
      text-transform: none;
    }}

    .footer-note {{
      margin-top: 14px;
      color: var(--muted);
      font-size: 11px;
      letter-spacing: 0.12em;
      text-transform: uppercase;
      display: flex;
      justify-content: space-between;
      flex-wrap: wrap;
      gap: 10px;
      border-top: 1px solid var(--line);
      padding-top: 12px;
    }}

    .hidden {{ display: none !important; }}

    @media (max-width: 980px) {{
      .title-row, .controls {{ grid-template-columns: 1fr; }}
      .chart-wrap {{ min-height: 0; }}
    }}

    @media (prefers-reduced-motion: reduce) {{
      *, *::before, *::after {{
        animation-duration: 0.001ms !important;
        animation-iteration-count: 1 !important;
        transition-duration: 0.001ms !important;
        scroll-behavior: auto !important;
      }}
    }}
  </style>
</head>
<body>
  <div class="shell">
    <div class="header-bar">
      <div class="brand">Smithsonian GVP archive / 2000–2025</div>
      <div class="status">Global eruption pulse</div>
    </div>

    <div class="title-row">
      <div class="hero">
        <div class="eyebrow">Smithsonian archive / 2000–2025</div>
        <div class="title-stack">
          <h1><span class="word-volcanic">Volcanic</span><span class="word-pulse">Pulse</span></h1>
          <div class="year-badge"><span>Current year</span><strong id="current-year-label">2025</strong></div>
        </div>
        <div class="subtitle">Scientific rhythm map of eruptions from 2000–2025, each point one confirmed eruption, VEI encoded in size and thermal scale.</div>
      </div>

      <div class="panel" aria-live="polite">
        <div class="metrics">
          <div class="metric">
            <span class="metric-label">Visible</span>
            <span id="visible-count" class="metric-value gold">{total_count}</span>
          </div>
          <div class="metric">
            <span class="metric-label">Unknown VEI</span>
            <span id="unknown-count" class="metric-value cyan">{unknown_count}</span>
          </div>
          <div class="metric">
            <span class="metric-label">Peak year</span>
            <span id="peak-year" class="metric-value orange">{max((year for year, count in y_counts.items() if count == max_year_count), default='—') if y_counts else '—'}</span>
          </div>
          <div class="metric">
            <span class="metric-label">Max in year</span>
            <span id="peak-count" class="metric-value red">{max_year_count}</span>
          </div>
        </div>
      </div>
    </div>

    <div class="chart-wrap">
      <div class="top-year-control">
        <div class="slider-wrap">
          <div class="slider-label">Year</div>
          <input id="top-year-slider" type="range" min="2000" max="2025" value="2025" step="1" />
          <div id="year-value-top" class="year-num">2025</div>
        </div>
      </div>
      <div class="hud" aria-hidden="true">
        <span class="corner tl"></span>
        <span class="corner tr"></span>
        <span class="corner bl"></span>
        <span class="corner br"></span>
      </div>
      <svg id="chart" viewBox="0 0 1200 1420" role="img" aria-label="Interactive volcanic eruption timeline"></svg>
      <div id="tooltip" class="tooltip" role="status" aria-live="polite"></div>
    </div>

    <div class="controls">
      <div class="control-panel">
        <div class="control-head">
          <span>Temporal view</span>
          <span id="status-text">Accumulated</span>
        </div>

        <div class="play-row">
          <button id="play-toggle" class="btn primary" type="button">Play</button>
          <button id="replay-btn" class="btn" type="button">Replay</button>
          <div class="speed-group" aria-label="Playback speed">
            <button type="button" data-speed="0.5">0.5x</button>
            <button type="button" data-speed="1" class="active">1x</button>
            <button type="button" data-speed="2">2x</button>
          </div>
        </div>

        <div class="slider-wrap">
          <div class="slider-label">Year</div>
          <input id="year-slider" type="range" min="2000" max="2025" value="2025" step="1" />
          <div id="year-value" class="year-num">2025</div>
        </div>

        <div class="mode-row" aria-label="Timeline mode">
          <label class="toggle"><input type="radio" name="mode" value="accumulate" checked /> Accumulate</label>
          <label class="toggle"><input type="radio" name="mode" value="single" /> Single year</label>
        </div>

        <div class="control-head" style="margin-top: 18px; margin-bottom: 10px;">
          <span>VEI filter</span>
          <span id="filter-count">{total_count} shown</span>
        </div>

        <div class="filters" id="filters">
          {''.join(f'<label class="filter-chip"><input type="checkbox" value="{label}" checked /><span class="dot" style="background:{FILTER_COLORS[label]}; border-color:{FILTER_COLORS[label]};"></span><span>{FILTER_LABELS[label]}</span></label>' for label in FILTER_ORDER)}
        </div>
      </div>

      <aside class="side-panel">
        <div class="detail-header">
          <span>Event detail</span>
          <span id="detail-tier">Selected</span>
        </div>
        <div id="detail-panel" class="detail-body">
          <div class="detail-name">No event selected</div>
          <ul class="detail-list">
            <li><strong>Year</strong> — none</li>
            <li><strong>VEI</strong> — none</li>
            <li><strong>Identifier</strong> — none</li>
          </ul>
        </div>
        <div class="detail-chip"><span class="dot"></span> High-VEI impact ring</div>
      </aside>
    </div>

    <div class="footer-note">
      <span>Source: Smithsonian GVP / Confirmed Holocene eruptions</span>
      <span>Offline self-contained HTML — SVG export — no geospatial map</span>
    </div>
  </div>

  <script>
    const DATA = {dataset_json};
    const YEAR_MIN = 2000;
    const YEAR_MAX = 2025;
    const FULL_VIEW = 1200;
    const FULL_HEIGHT = 980;
    const CENTER_X = FULL_VIEW / 2;
    const CENTER_Y = FULL_HEIGHT / 2;
    const CHART_TOP = 150;
    const CHART_BOTTOM = 820;
    const YEAR_COLUMN_GAP = 40;
    const reduceMotion = window.matchMedia('(prefers-reduced-motion: reduce)').matches;

    const state = {{
      filters: {{ low: true, medium: true, high: true, unknown: true }},
      mode: 'accumulate',
      year: YEAR_MAX,
      playTimer: null,
      playing: false,
      speed: 1,
      selectedId: null,
      currentVisible: []
    }};

    const chartRefs = {{
      svg: null,
      background: null,
      labelLayer: null,
      pointLayer: null,
      yearLabels: new Map(),
      circleMap: new Map(),
      haloMap: new Map(),
      visibleIds: new Set()
    }};

    function isUnknown(vei) {{
      return vei === null || typeof vei === 'undefined';
    }}

    function clamp(value, min, max) {{
      return Math.min(Math.max(value, min), max);
    }}

    function getGroupFromVEI(vei) {{
      if (isUnknown(vei)) return 'unknown';
      if (vei <= 1) return 'low';
      if (vei <= 3) return 'medium';
      return 'high';
    }}

    function getYearColumnX(year) {{
      return 120 + (year - YEAR_MIN) * YEAR_COLUMN_GAP;
    }}

    function getLayoutSettings() {{
      const maxOrder = Math.max(1, ...DATA.map((item) => item.order));
      const pointSpacing = Math.max(13.5, Math.min(16, (CHART_BOTTOM - CHART_TOP) / Math.max(1, maxOrder - 1)));
      return {{ pointSpacing, chartTop: CHART_TOP, chartBottom: CHART_BOTTOM }};
    }}

    function getPointPosition(item) {{
      const {{ pointSpacing, chartTop }} = getLayoutSettings();
      const x = getYearColumnX(item.year);
      const y = chartTop + (item.order - 1) * pointSpacing;
      return {{ x, y, pointSpacing }};
    }}

    function getPointRadius(item) {{
      if (isUnknown(item.vei)) return 4.2;
      if (item.group === 'low') return 4.1;
      if (item.group === 'medium') return 5.8;
      return 8.2;
    }}

    function getHaloRadius(item) {{
      if (item.group !== 'high') return 0;
      return clamp((item.vei || 4) * 2.6 + 13, 13, 26);
    }}

    function getVisibleEntries() {{
      return DATA.filter((entry) => {{
        if (!state.filters[entry.group]) return false;
        if (state.mode === 'single') return entry.year === state.year;
        return entry.year <= state.year;
      }});
    }}

    function setYear(nextYear, options = {{}}) {{
      const safeYear = clamp(Number(nextYear), YEAR_MIN, YEAR_MAX);
      const {{ manual = false }} = options;
      if (manual) {{
        clearTimeout(state.playTimer);
        state.playing = false;
        document.getElementById('play-toggle').textContent = 'Play';
      }}
      state.year = safeYear;
      updateYearLabel();
      updateFilterCount();
      renderChart();
    }}

    function updateYearLabel() {{
      const yearText = document.getElementById('year-value');
      const topYearText = document.getElementById('year-value-top');
      const bottomSlider = document.getElementById('year-slider');
      const topSlider = document.getElementById('top-year-slider');

      if (yearText) yearText.textContent = String(state.year);
      if (topYearText) topYearText.textContent = String(state.year);
      if (bottomSlider) bottomSlider.value = String(state.year);
      if (topSlider) topSlider.value = String(state.year);
      const statusText = document.getElementById('status-text');
      if (statusText) statusText.textContent = state.mode === 'single' ? 'Single year' : 'Accumulated';
      const currentYearLabel = document.getElementById('current-year-label');
      if (currentYearLabel) currentYearLabel.textContent = String(state.year);
    }}

    function updateFilterCount() {{
      const visible = getVisibleEntries();
      const visibleCount = document.getElementById('visible-count');
      const filterCount = document.getElementById('filter-count');
      const unknownCount = document.getElementById('unknown-count');
      if (visibleCount) visibleCount.textContent = String(visible.length);
      if (filterCount) filterCount.textContent = `${{visible.length}} shown`;
      if (unknownCount) unknownCount.textContent = String(visible.filter(item => isUnknown(item.vei)).length);
      state.currentVisible = visible;
    }}

    function renderDetailPanel() {{
      const panel = document.getElementById('detail-panel');
      const selected = DATA.find(item => item.id === state.selectedId) || null;
      if (!selected) {{
        panel.innerHTML = `
          <div class="detail-name">No event selected</div>
          <ul class="detail-list">
            <li><strong>Year</strong> — none</li>
            <li><strong>VEI</strong> — none</li>
            <li><strong>Identifier</strong> — none</li>
          </ul>
        `;
        return;
      }}
      const tier = selected.group === 'high' ? 'High-VEI event' : (selected.group === 'unknown' ? 'Unknown VEI' : 'Temporal event');
      const eruptionText = selected.eruption_number ? `#${{selected.eruption_number}}` : 'Not recorded';
      panel.innerHTML = `
        <div class="detail-name">${{selected.volcano}}</div>
        <ul class="detail-list">
          <li><strong>Start year</strong> — ${{selected.year}}</li>
          <li><strong>VEI</strong> — ${{selected.vei_label}}</li>
          <li><strong>Start date</strong> — ${{selected.date}}</li>
          <li><strong>Eruption number</strong> — ${{eruptionText}}</li>
        </ul>
      `;
      document.getElementById('detail-tier').textContent = tier;
    }}

    function showTooltip(event, text) {{
      const tooltip = document.getElementById('tooltip');
      tooltip.innerHTML = text;
      tooltip.classList.add('visible');
      const rect = document.querySelector('.chart-wrap').getBoundingClientRect();
      tooltip.style.left = event.clientX - rect.left + 16 + 'px';
      tooltip.style.top = event.clientY - rect.top + 10 + 'px';
    }}

    function hideTooltip() {{
      document.getElementById('tooltip').classList.remove('visible');
    }}

    function ensureChartStructure() {{
      const svg = document.getElementById('chart');
      chartRefs.svg = svg;
      svg.setAttribute('viewBox', `0 0 ${{FULL_VIEW}} ${{FULL_HEIGHT}}`);

      if (!chartRefs.background) {{
        const background = document.createElementNS('http://www.w3.org/2000/svg', 'g');
        background.dataset.layer = 'background';
        const bgMarkup = `
          <rect x="0" y="0" width="${{FULL_VIEW}}" height="${{FULL_HEIGHT}}" fill="rgba(9,11,16,0.15)" />
          <g opacity="0.76">
            <circle cx="610" cy="490" r="120" fill="none" stroke="rgba(56,214,230,0.12)" stroke-width="1" stroke-dasharray="3 8" />
            <circle cx="610" cy="490" r="180" fill="none" stroke="rgba(56,214,230,0.10)" stroke-width="1" stroke-dasharray="2 10" />
            <circle cx="610" cy="490" r="240" fill="none" stroke="rgba(245,199,110,0.10)" stroke-width="1" stroke-dasharray="4 8" />
            <circle cx="610" cy="490" r="320" fill="none" stroke="rgba(56,214,230,0.07)" stroke-width="1" stroke-dasharray="2 12" />
            <circle cx="610" cy="490" r="390" fill="none" stroke="rgba(245,199,110,0.05)" stroke-width="1" stroke-dasharray="2 14" />
          </g>
          <g opacity="0.38">
            <path d="M 610 300 A 200 200 0 0 1 810 490" fill="none" stroke="rgba(56,214,230,0.22)" stroke-width="1" />
            <path d="M 610 300 A 200 200 0 0 0 410 490" fill="none" stroke="rgba(245,199,110,0.16)" stroke-width="1" />
            <path d="M 610 680 A 200 200 0 0 1 810 490" fill="none" stroke="rgba(245,199,110,0.14)" stroke-width="1" />
            <path d="M 610 680 A 200 200 0 0 0 410 490" fill="none" stroke="rgba(56,214,230,0.14)" stroke-width="1" />
          </g>
        `;
        background.innerHTML = bgMarkup;
        svg.appendChild(background);
        chartRefs.background = background;
      }}

      if (!chartRefs.labelLayer) {{
        const labelLayer = document.createElementNS('http://www.w3.org/2000/svg', 'g');
        labelLayer.dataset.layer = 'labels';
        svg.appendChild(labelLayer);
        chartRefs.labelLayer = labelLayer;
        for (let year = YEAR_MIN; year <= YEAR_MAX; year += 1) {{
          const x = getYearColumnX(year);
          const label = document.createElementNS('http://www.w3.org/2000/svg', 'text');
          label.setAttribute('x', String(x));
          label.setAttribute('y', String(CHART_TOP - 64));
          label.setAttribute('fill', 'rgba(196,204,214,0.8)');
          label.setAttribute('font-size', '10');
          label.setAttribute('font-weight', '700');
          label.setAttribute('text-anchor', 'middle');
          label.setAttribute('letter-spacing', '2');
          label.textContent = String(year);
          labelLayer.appendChild(label);
          chartRefs.yearLabels.set(year, label);
        }}
      }}

      if (!chartRefs.pointLayer) {{
        const pointLayer = document.createElementNS('http://www.w3.org/2000/svg', 'g');
        pointLayer.dataset.layer = 'points';
        svg.appendChild(pointLayer);
        chartRefs.pointLayer = pointLayer;
      }}

      const currentX = getYearColumnX(state.year);
      const rect = document.querySelector('[data-year-band]') || document.createElementNS('http://www.w3.org/2000/svg', 'rect');
      rect.setAttribute('data-year-band', 'true');
      rect.setAttribute('x', String(currentX - 24));
      rect.setAttribute('y', String(CHART_TOP - 26));
      rect.setAttribute('width', '48');
      rect.setAttribute('height', String(CHART_BOTTOM - CHART_TOP + 90));
      rect.setAttribute('rx', '18');
      rect.setAttribute('fill', 'rgba(245,199,110,0.06)');
      rect.setAttribute('stroke', 'rgba(245,199,110,0.18)');
      rect.setAttribute('stroke-width', '1');
      if (!rect.parentNode) {{
        chartRefs.background.appendChild(rect);
      }}
    }}

    function updateYearLabels() {{
      const currentYear = clamp(state.year, YEAR_MIN, YEAR_MAX);
      chartRefs.yearLabels.forEach((label, year) => {{
        label.setAttribute('fill', year === currentYear ? 'rgba(245,199,110,0.9)' : 'rgba(196,204,214,0.8)');
        label.setAttribute('font-weight', year === currentYear ? '800' : '700');
      }});
      const rect = chartRefs.background.querySelector('[data-year-band]');
      if (rect) {{
        const currentX = getYearColumnX(currentYear);
        rect.setAttribute('x', String(currentX - 24));
      }}
    }}

    function createCircleNode(item) {{
      const group = item.group;
      const color = group === 'low' ? '#F5C76E' : group === 'medium' ? '#FF8C42' : group === 'high' ? '#FF593B' : 'rgba(0,0,0,0)';
      const radius = getPointRadius(item);
      const pos = getPointPosition(item);

      const halo = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
      const haloRadius = getHaloRadius(item);
      halo.setAttribute('cx', String(pos.x));
      halo.setAttribute('cy', String(pos.y));
      halo.setAttribute('r', String(haloRadius));
      halo.setAttribute('fill', 'none');
      halo.setAttribute('stroke', group === 'high' ? 'rgba(255, 130, 78, 0.72)' : 'transparent');
      halo.setAttribute('stroke-width', group === 'high' ? String(0.9 + (item.vei || 4) * 0.3) : '0');
      halo.setAttribute('opacity', group === 'high' ? String(0.38 + (item.vei || 4) * 0.09) : '0');
      halo.setAttribute('filter', group === 'high' ? 'drop-shadow(0 0 10px rgba(255, 129, 77, 0.85))' : 'none');
      halo.style.pointerEvents = 'none';
      halo.style.transition = 'opacity 0.22s ease';

      const circle = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
      circle.setAttribute('cx', String(pos.x));
      circle.setAttribute('cy', String(pos.y));
      circle.setAttribute('r', String(radius));
      circle.setAttribute('fill', group === 'unknown' ? 'rgba(0,0,0,0)' : color);
      circle.setAttribute('stroke', group === 'unknown' ? '#8C97A9' : color);
      circle.setAttribute('stroke-width', group === 'unknown' ? '1.5' : '0.9');
      circle.setAttribute('opacity', '0.9');
      circle.style.cursor = 'pointer';
      circle.style.transformOrigin = 'center';
      circle.style.transform = 'scale(0.9)';
      circle.style.transition = 'opacity 0.22s ease, transform 0.22s ease';
      circle.dataset.id = item.id;
      circle.dataset.group = item.group;
      circle.dataset.year = String(item.year);
      circle.dataset.vei = item.vei_label;
      circle.dataset.volcano = item.volcano;
      circle.dataset.date = item.date;
      circle.dataset.number = item.eruption_number || 'n/a';

      circle.addEventListener('mouseenter', (event) => {{
        showTooltip(event, `
          <div class="tooltip-head">${{escape(item.volcano)}}</div>
          <div class="tooltip-row"><span>Year</span><strong>${{item.year}}</strong></div>
          <div class="tooltip-row"><span>VEI</span><strong>${{item.vei_label}}</strong></div>
          <div class="tooltip-row"><span>Eruption</span><strong>${{item.eruption_number ? '#' + escape(item.eruption_number) : 'Not recorded'}}</strong></div>
        `);
      }});
      circle.addEventListener('mousemove', (event) => {{
        const rect = document.querySelector('.chart-wrap').getBoundingClientRect();
        const tooltip = document.getElementById('tooltip');
        tooltip.style.left = event.clientX - rect.left + 16 + 'px';
        tooltip.style.top = event.clientY - rect.top + 10 + 'px';
      }});
      circle.addEventListener('mouseleave', hideTooltip);
      circle.addEventListener('click', () => {{
        state.selectedId = item.id;
        renderDetailPanel();
        const point = getPointPosition(item);
        if (item.group === 'high' || item.group === 'medium') {{
          showImpactPulse(point.x, point.y, item.group);
          spawnImpactParticles(point.x, point.y, item.group);
        }}
      }});

      chartRefs.pointLayer.appendChild(halo);
      chartRefs.pointLayer.appendChild(circle);
      chartRefs.circleMap.set(item.id, circle);
      chartRefs.haloMap.set(item.id, halo);

      return {{ circle, halo }};
    }}

    function setPointState(item, circle, halo) {{
      const pos = getPointPosition(item);
      const radius = getPointRadius(item);
      const activeYear = item.year === state.year;
      const visible = item.group && state.filters[item.group] && (state.mode !== 'single' ? item.year <= state.year : item.year === state.year);
      circle.setAttribute('cx', String(pos.x));
      circle.setAttribute('cy', String(pos.y));
      circle.setAttribute('r', String(radius));
      circle.setAttribute('fill', item.group === 'unknown' ? 'rgba(0,0,0,0)' : item.group === 'low' ? '#F5C76E' : item.group === 'medium' ? '#FF8C42' : '#FF593B');
      circle.setAttribute('stroke', item.group === 'unknown' ? '#8C97A9' : activeYear ? '#EAFBFF' : item.group === 'low' ? '#F5C76E' : item.group === 'medium' ? '#FF8C42' : '#FF593B');
      circle.setAttribute('stroke-width', item.group === 'unknown' ? '1.5' : activeYear ? '1.6' : '0.9');
      circle.style.filter = activeYear ? 'drop-shadow(0 0 12px rgba(255, 204, 124, 0.92)) drop-shadow(0 0 22px rgba(90, 198, 255, 0.65))' : 'none';
      circle.style.opacity = visible ? '1' : '0';
      circle.style.transform = visible ? 'scale(1)' : 'scale(0.82)';
      circle.style.pointerEvents = visible ? 'auto' : 'none';
      circle.style.display = visible ? 'block' : 'none';
      if (halo) {{
        halo.setAttribute('cx', String(pos.x));
        halo.setAttribute('cy', String(pos.y));
        halo.setAttribute('r', String(getHaloRadius(item)));
        halo.setAttribute('opacity', item.group === 'high' && visible ? String(0.38 + (item.vei || 4) * 0.09) : '0');
        halo.style.display = item.group === 'high' && visible ? 'block' : 'none';
      }}
    }}

    function syncVisiblePoints() {{
      const visibleEntries = getVisibleEntries();
      const desiredIds = new Set(visibleEntries.map(item => item.id));
      const previousIds = new Set(chartRefs.visibleIds || []);

      visibleEntries.forEach((item) => {{
        let circleNode = chartRefs.circleMap.get(item.id);
        if (!circleNode) {{
          const created = createCircleNode(item);
          circleNode = created.circle;
        }}
        const halo = chartRefs.haloMap.get(item.id);
        const shouldAnimateEnter = !previousIds.has(item.id);
        const shouldAnimateExit = previousIds.has(item.id) && !desiredIds.has(item.id);
        if (shouldAnimateExit) return;
        setPointState(item, circleNode, halo);
        if (shouldAnimateEnter) {{
          circleNode.style.opacity = '0';
          circleNode.style.transform = 'scale(0.9)';
          requestAnimationFrame(() => {{
            circleNode.style.opacity = '1';
            circleNode.style.transform = 'scale(1)';
          }});
        }}
      }});

      for (const id of previousIds) {{
        if (!desiredIds.has(id)) {{
          const circle = chartRefs.circleMap.get(id);
          if (circle) {{
            circle.style.opacity = '0';
            circle.style.transform = 'scale(0.82)';
            circle.style.pointerEvents = 'none';
            circle.style.display = 'none';
            const halo = chartRefs.haloMap.get(id);
            if (halo) {{
              halo.style.display = 'none';
            }}
          }}
        }}
      }}

      chartRefs.visibleIds = desiredIds;
    }}

    function spawnImpactParticles(x, y, group) {{
      if (reduceMotion || group !== 'high') return;
      const svg = document.getElementById('chart');
      const count = 10 + Math.floor(Math.random() * 5);
      for (let index = 0; index < count; index += 1) {{
        const particle = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
        const angle = (Math.PI * 2 * index) / count + (Math.random() - 0.5) * 0.8;
        const distance = 10 + Math.random() * 34;
        const px = x + Math.cos(angle) * distance;
        const py = y + Math.sin(angle) * distance;
        particle.setAttribute('cx', String(x));
        particle.setAttribute('cy', String(y));
        particle.setAttribute('r', String(2 + Math.random() * 2.5));
        particle.setAttribute('fill', index % 2 === 0 ? '#F5C76E' : '#FF8C42');
        particle.setAttribute('opacity', '0.9');
        particle.style.transition = 'opacity 0.24s ease';
        svg.appendChild(particle);
        const duration = 560 + Math.random() * 320;
        requestAnimationFrame(() => {{
          particle.setAttribute('cx', String(px));
          particle.setAttribute('cy', String(py));
          particle.setAttribute('opacity', '0.15');
        }});
        setTimeout(() => {{ particle.remove(); }}, duration);
      }}
    }}

    function showImpactPulse(x, y, group) {{
      if (reduceMotion) return;
      if (group !== 'high' && group !== 'medium') return;
      const svg = document.getElementById('chart');
      const ring = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
      const ringColor = group === 'high' ? '#FF593B' : '#FF8C42';
      ring.setAttribute('cx', String(x));
      ring.setAttribute('cy', String(y));
      ring.setAttribute('r', '8');
      ring.setAttribute('fill', 'none');
      ring.setAttribute('stroke', ringColor);
      ring.setAttribute('stroke-width', '2');
      ring.setAttribute('opacity', '0.9');
      ring.style.transition = 'opacity 0.22s ease, r 0.22s ease';
      svg.appendChild(ring);
      const burst = [0.6, 1.0, 1.4, 1.8].map((scale, idx) => {{
        const extra = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
        extra.setAttribute('cx', String(x));
        extra.setAttribute('cy', String(y));
        extra.setAttribute('r', String(8 + idx * 4));
        extra.setAttribute('fill', 'none');
        extra.setAttribute('stroke', ringColor);
        extra.setAttribute('stroke-width', '1');
        extra.setAttribute('opacity', String(0.7 - idx * 0.18));
        extra.style.transition = 'opacity 0.22s ease';
        svg.appendChild(extra);
        return extra;
      }});
      const pulses = [ring, ...burst];
      pulses.forEach((element, index) => {{
        const start = index === 0 ? 0 : index * 120;
        setTimeout(() => {{
          element.setAttribute('r', String(10 + index * 8));
          element.setAttribute('opacity', '0');
        }}, start);
      }});
      setTimeout(() => {{
        pulses.forEach((element) => element.remove());
      }}, 950);
    }}

    function renderChart() {{
      ensureChartStructure();
      const currentYear = clamp(state.year, YEAR_MIN, YEAR_MAX);
      updateYearLabels();
      syncVisiblePoints();

      const activeVisible = getVisibleEntries();
      activeVisible.forEach((item) => {{
        const circle = chartRefs.circleMap.get(item.id);
        if (!circle) return;
        const isCurrent = item.year === currentYear;
        circle.setAttribute('stroke', item.group === 'unknown' ? '#8C97A9' : isCurrent ? '#EAFBFF' : item.group === 'low' ? '#F5C76E' : item.group === 'medium' ? '#FF8C42' : '#FF593B');
        circle.setAttribute('stroke-width', item.group === 'unknown' ? '1.5' : isCurrent ? '1.6' : '0.9');
        circle.style.filter = isCurrent ? 'drop-shadow(0 0 12px rgba(255, 204, 124, 0.92)) drop-shadow(0 0 22px rgba(90, 198, 255, 0.65))' : 'none';
      }});

      if (state.selectedId) {{
        const selected = DATA.find((item) => item.id === state.selectedId);
        if (selected) {{
          const selectedCircle = chartRefs.circleMap.get(selected.id);
          if (selectedCircle) {{
            const pos = getPointPosition(selected);
            const ring = document.createElementNS('http://www.w3.org/2000/svg', 'circle');
            ring.setAttribute('cx', String(pos.x));
            ring.setAttribute('cy', String(pos.y));
            ring.setAttribute('r', '14');
            ring.setAttribute('fill', 'none');
            ring.setAttribute('stroke', '#F5C76E');
            ring.setAttribute('stroke-width', '1.5');
            ring.setAttribute('opacity', '0.9');
            ring.setAttribute('stroke-dasharray', '3 6');
            if (!ring.dataset.selected) {{
              ring.dataset.selected = 'true';
              chartRefs.pointLayer.appendChild(ring);
            }}
            ring.setAttribute('cx', String(pos.x));
            ring.setAttribute('cy', String(pos.y));
          }}
        }}
      }}
    }}

    function tickPlayback() {{
      if (!state.playing) return;
      const nextYear = state.year + 1;
      if (nextYear > YEAR_MAX) {{
        state.year = YEAR_MAX;
        state.playing = false;
        document.getElementById('play-toggle').textContent = 'Play';
        updateYearLabel();
        updateFilterCount();
        renderChart();
        return;
      }}
      state.year = nextYear;
      updateYearLabel();
      updateFilterCount();
      renderChart();
      const interval = reduceMotion ? 180 : 700 / state.speed;
      state.playTimer = setTimeout(tickPlayback, interval);
    }}

    function handleFilterChange(event) {{
      const {{ value }} = event.target;
      state.filters[value] = event.target.checked;
      updateFilterCount();
      renderChart();
      renderDetailPanel();
    }}

    function bindControls() {{
      const syncYear = (value, options = {{}}) => setYear(value, {{ ...options, manual: true }});

      document.getElementById('year-slider').addEventListener('input', (event) => syncYear(event.target.value));
      document.getElementById('top-year-slider').addEventListener('input', (event) => syncYear(event.target.value));

      document.querySelectorAll('input[name="mode"]').forEach((input) => {{
        input.addEventListener('change', () => {{
          state.mode = input.value;
          if (state.mode === 'single') {{
            state.year = clamp(state.year, YEAR_MIN, YEAR_MAX);
          }} else {{
            state.year = YEAR_MAX;
          }}
          updateYearLabel();
          updateFilterCount();
          renderChart();
        }});
      }});

      document.getElementById('play-toggle').addEventListener('click', () => {{
        state.playing = !state.playing;
        if (state.playing) {{
          if (state.mode !== 'accumulate') {{
            state.mode = 'accumulate';
            document.querySelector('input[name="mode"][value="accumulate"]').checked = true;
          }}
          state.year = YEAR_MIN;
          updateYearLabel();
          updateFilterCount();
          renderChart();
          const interval = reduceMotion ? 160 : 700 / state.speed;
          state.playTimer = setTimeout(tickPlayback, interval);
        }} else {{
          clearTimeout(state.playTimer);
        }}
        document.getElementById('play-toggle').textContent = state.playing ? 'Pause' : 'Play';
      }});

      document.getElementById('replay-btn').addEventListener('click', () => {{
        clearTimeout(state.playTimer);
        state.playing = false;
        state.mode = 'accumulate';
        state.year = YEAR_MIN;
        document.querySelector('input[name="mode"][value="accumulate"]').checked = true;
        document.getElementById('play-toggle').textContent = 'Play';
        updateYearLabel();
        updateFilterCount();
        renderChart();
      }});

      document.querySelectorAll('[data-speed]').forEach((button) => {{
        button.addEventListener('click', () => {{
          state.speed = Number(button.dataset.speed);
          document.querySelectorAll('[data-speed]').forEach((item) => item.classList.toggle('active', item === button));
          if (state.playing) {{
            clearTimeout(state.playTimer);
            const interval = reduceMotion ? 160 : 700 / state.speed;
            state.playTimer = setTimeout(tickPlayback, interval);
          }}
        }});
      }});

      document.querySelectorAll('#filters input[type="checkbox"]').forEach((checkbox) => {{
        checkbox.addEventListener('change', handleFilterChange);
      }});

      document.getElementById('tooltip').addEventListener('mouseleave', hideTooltip);
      document.addEventListener('keydown', (event) => {{
        if (event.target && ['INPUT', 'TEXTAREA', 'SELECT'].includes(event.target.tagName)) return;
        if (event.key === 'ArrowLeft') {{
          event.preventDefault();
          setYear(state.year - 1, {{ manual: true }});
        }}
        if (event.key === 'ArrowRight') {{
          event.preventDefault();
          setYear(state.year + 1, {{ manual: true }});
        }}
      }});
    }}

    function exportSvg() {{
      const svg = document.getElementById('chart');
      const serializer = new XMLSerializer();
      const source = serializer.serializeToString(svg);
      const blob = new Blob([source], {{ type: 'image/svg+xml;charset=utf-8' }});
      const url = URL.createObjectURL(blob);
      const a = document.createElement('a');
      a.href = url;
      a.download = 'volcanic-pulse.svg';
      document.body.appendChild(a);
      a.click();
      a.remove();
      URL.revokeObjectURL(url);
    }}

    function exportPng() {{
      const svg = document.getElementById('chart');
      const serializer = new XMLSerializer();
      const source = serializer.serializeToString(svg);
      const svgBlob = new Blob([source], {{ type: 'image/svg+xml;charset=utf-8' }});
      const url = URL.createObjectURL(svgBlob);
      const img = new Image();
      img.onload = () => {{
        const canvas = document.createElement('canvas');
        const viewBox = svg.viewBox.baseVal;
        const scale = 2;
        canvas.width = Math.round(viewBox.width * scale);
        canvas.height = Math.round(viewBox.height * scale);
        const ctx = canvas.getContext('2d');
        ctx.fillStyle = '#090B10';
        ctx.fillRect(0, 0, canvas.width, canvas.height);
        ctx.drawImage(img, 0, 0, canvas.width, canvas.height);
        const pngUrl = canvas.toDataURL('image/png');
        const a = document.createElement('a');
        a.href = pngUrl;
        a.download = 'volcanic-pulse.png';
        document.body.appendChild(a);
        a.click();
        a.remove();
        URL.revokeObjectURL(url);
      }};
      img.src = url;
    }}

    function addExportButtons() {{
      const controls = document.querySelector('.play-row');
      const svgBtn = document.createElement('button');
      svgBtn.type = 'button';
      svgBtn.className = 'btn';
      svgBtn.textContent = 'Export SVG';
      svgBtn.addEventListener('click', exportSvg);
      controls.appendChild(svgBtn);

      const pngBtn = document.createElement('button');
      pngBtn.type = 'button';
      pngBtn.className = 'btn';
      pngBtn.textContent = 'Export PNG';
      pngBtn.addEventListener('click', exportPng);
      controls.appendChild(pngBtn);
    }}

    function init() {{
      bindControls();
      addExportButtons();
      updateYearLabel();
      updateFilterCount();
      renderDetailPanel();
      ensureChartStructure();
      renderChart();
    }}

    init();
  </script>
</body>
</html>
'''
    return page


def main():
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Missing source workbook: {DATA_FILE}")

    rows, parse_errors = load_eruption_rows(DATA_FILE)
    print(f"source: {DATA_FILE.name}")
    print(f"records_in_range: {len(rows)}")
    print(f"unknown_vei: {sum(1 for item in rows if item['vei'] is None)}")
    print(f"parse_errors: {parse_errors}")

    html = build_page(rows, parse_errors)
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(html, encoding="utf-8")
    print(f"written: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()


